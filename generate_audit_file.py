import os
import sys
import datetime
import django
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Setup Django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ats.settings')
django.setup()

from apps.candidates.models import Candidate, JobApplication, ApplicationStageState
from apps.jobs.models import JobOpportunity, JobOpportunityStage

print("Starting generation of comprehensive audit & human-review workbook...")

# Paths
INPUT_EXCEL_PATH = '/Users/omidsalehi/Downloads/ورود اطلاعات.xlsx'
OUTPUT_EXCEL_WORKSPACE = '/Users/omidsalehi/Downloads/Payjoo-ATS-Demo-main/بررسی_و_ورود_غایبین_و_سوابق_فولاد_مبارکه.xlsx'
OUTPUT_EXCEL_DOWNLOADS = '/Users/omidsalehi/Downloads/بررسی_و_ورود_غایبین_و_سوابق_فولاد_مبارکه.xlsx'

# Load input Excel data
wb_in = openpyxl.load_workbook(INPUT_EXCEL_PATH, data_only=True)

# Helper to normalize national ID and codes
def norm_code(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s

def norm_nid(val):
    s = norm_code(val)
    if s.isdigit():
        return s.zfill(10)
    return s

# Pre-fetch DB caches
print("Pre-fetching DB data...")
candidates_by_nid = {c.national_id: c for c in Candidate.objects.filter(is_deleted=False)}
jobs_by_code = {j.code: j for j in JobOpportunity.objects.filter(is_deleted=False)}

# Build application & stage states map: (job_code, national_id, stage_type) -> ApplicationStageState
stage_type_map = {
    'SCREENING': 'غربالگری',
    'EXAM': 'کتبی',
    'SKILL_TEST': 'مهارتی',
    'INTERVIEW': 'مصاحبه',
    'ASSESSMENT': 'کانون'
}

type_to_stage_name = {
    'EXAM': 'آزمون کتبی',
    'SKILL_TEST': 'آزمون مهارتی',
    'INTERVIEW': 'مصاحبه تخصصی',
    'ASSESSMENT': 'کانون ارزیابی'
}

print("Mapping stage states from DB...")
states_cache = {}
for ss in ApplicationStageState.objects.filter(is_deleted=False).select_related('application__job', 'application__candidate', 'stage'):
    app = ss.application
    if app and app.job and app.candidate and ss.stage:
        j_code = app.job.code
        nid = app.candidate.national_id
        st_type = ss.stage.stage_type
        states_cache[(j_code, nid, st_type)] = ss

# Read 'جدول وضعیت' from Excel
jobs_excel_info = {}
if 'جدول وضعیت' in wb_in.sheetnames:
    ws_jobs = wb_in['جدول وضعیت']
    rows = list(ws_jobs.iter_rows(values_only=True))
    if rows:
        headers = [str(h).strip() if h is not None else '' for h in rows[0]]
        code_idx = headers.index('کد') if 'کد' in headers else 2
        title_idx = headers.index('عنوان پست') if 'عنوان پست' in headers else 6
        unit_idx = headers.index('واحد متقاضی') if 'واحد متقاضی' in headers else 5
        cat_idx = headers.index('رده شغلی') if 'رده شغلی' in headers else 7
        req_idx = headers.index('تعداد مورد نیاز') if 'تعداد مورد نیاز' in headers else 8

        for r in rows[1:]:
            if not r or r[code_idx] is None:
                continue
            c = norm_code(r[code_idx])
            jobs_excel_info[c] = {
                'title': str(r[title_idx] or '').strip(),
                'unit': str(r[unit_idx] or '').strip(),
                'category': str(r[cat_idx] or '').strip(),
                'required': r[req_idx]
            }

# Read stages results from Excel
stage_sheets_info = [
    ('کتبی', 'EXAM', 'ScoreW', 'Result1'),
    ('مهارتی', 'SKILL_TEST', 'ScoreS', 'Result2'),
    ('مصاحبه', 'INTERVIEW', 'ScoreI', 'Result3'),
    ('کانون', 'ASSESSMENT', 'ScoreAC', 'Result4'),
]

excel_stage_data = {} # (sheet_name) -> list of dicts
all_absentees = []
critical_discrepancies = []

for sheet_name, st_type, score_col, res_col in stage_sheets_info:
    ws = wb_in[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    exam_idx = headers.index('ExamCode')
    nat_idx = headers.index('NationCode')
    score_idx = headers.index(score_col) if score_col in headers else None
    res_idx = [i for i, h in enumerate(headers) if 'Result' in h][0]

    records = []
    for r in rows[1:]:
        if not r or r[nat_idx] is None:
            continue
        c_exam = norm_code(r[exam_idx])
        c_nat = norm_nid(r[nat_idx])
        c_score = r[score_idx] if score_idx is not None else 0.0
        c_res = str(r[res_idx]).strip() if r[res_idx] is not None else ''

        # Clean score
        try:
            val_score = float(str(c_score).strip()) if c_score is not None else 0.0
        except (ValueError, TypeError):
            val_score = 0.0

        # DB match
        ss = states_cache.get((c_exam, c_nat, st_type))
        cand = candidates_by_nid.get(c_nat)
        job = jobs_by_code.get(c_exam)
        j_info = jobs_excel_info.get(c_exam, {})

        rec = {
            'exam_code': c_exam,
            'national_id': c_nat,
            'score': val_score,
            'result_raw': c_res,
            'stage_type': st_type,
            'sheet_name': sheet_name,
            'candidate': cand,
            'job': job,
            'job_info': j_info,
            'stage_state': ss
        }
        records.append(rec)

        is_absent = 'غایب' in c_res
        if is_absent:
            all_absentees.append(rec)
            # Check discrepancy
            if ss and ss.status == 'COMPLETED':
                critical_discrepancies.append({
                    'rec': rec,
                    'type': 'تناقض بحرانی: در اکسل غایب است ولی در سامانه قبولی (COMPLETED) ثبت شده!',
                    'action': 'بررسی کارشناس: آیا نمره سامانه اشتباه است یا فرد حضور داشته؟ در صورت غیبت باید به غایب (ABSENT) اصلاح شود.'
                })
            elif not ss:
                reason = "فرصت شغلی در سامانه تعریف نشده است" if not job else ("پرونده یا مرحله متقاضی یافت نشد" if not cand else "مرحله متناظر در پرونده وجود ندارد")
                critical_discrepancies.append({
                    'rec': rec,
                    'type': f'عدم تطبیق سیستمی: {reason}',
                    'action': 'ایجاد شغل/پرونده و ثبت وضعیت غایب پس از تایید کارشناس'
                })

    excel_stage_data[sheet_name] = records

print(f"Extracted {len(all_absentees)} absentees in total.")
print(f"Identified {len(critical_discrepancies)} discrepancies / items requiring human decision.")

# Now create Master Excel Workbook
wb_out = openpyxl.Workbook()
# remove default sheet
default_sheet = wb_out.active

# Styling constants
NAVY_HEADER = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
STEEL_BLUE_SUB = PatternFill(start_color="335C8D", end_color="335C8D", fill_type="solid")
GREEN_PASS = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
RED_FAIL = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
GRAY_ABSENT = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")
AMBER_WARNING = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

FONT_HEADER = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
FONT_TITLE = Font(name="Tahoma", size=13, bold=True, color="1B365D")
FONT_SUBTITLE = Font(name="Tahoma", size=10, color="555555")
FONT_BOLD = Font(name="Tahoma", size=9, bold=True)
FONT_NORMAL = Font(name="Tahoma", size=9)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")

THIN_BORDER = Border(
    left=Side(style='thin', color='D0D7DE'),
    right=Side(style='thin', color='D0D7DE'),
    top=Side(style='thin', color='D0D7DE'),
    bottom=Side(style='thin', color='D0D7DE')
)

def format_sheet_view(ws):
    ws.views.sheetView[0].rightToLeft = True

# -------------------------------------------------------------
# SHEET 1: داشبورد_مدیریتی_و_راهنما
# -------------------------------------------------------------
ws1 = wb_out.create_sheet(title="داشبورد_مدیریتی_و_راهنما")
format_sheet_view(ws1)

ws1.merge_cells("A1:G1")
ws1["A1"] = "شرکت فولاد مبارکه اصفهان - معاونت سرمایه انسانی و سازماندهی"
ws1["A1"].font = FONT_TITLE
ws1["A1"].alignment = ALIGN_CENTER

ws1.merge_cells("A2:G2")
ws1["A2"] = "گزارش ممیزی سوابق تاریخی آزمون‌ها و دستورالعمل کنترل انسانی ورود اطلاعات غایبین به سامانه ATS"
ws1["A2"].font = Font(name="Tahoma", size=11, bold=True, color="335C8D")
ws1["A2"].alignment = ALIGN_CENTER

ws1.merge_cells("A3:G3")
ws1["A3"] = f"تاریخ گزارش: {datetime.date.today().strftime('%Y/%m/%d')} | نسخه فایل: ۲.۰ (مجهز به وضعیت تفکیکی غایب / ABSENT) | تهیه شده توسط: کارشناس تأمین سرمایه انسانی و راهبر سامانه"
ws1["A3"].font = FONT_SUBTITLE
ws1["A3"].alignment = ALIGN_CENTER

# Table 1: Stats Summary
ws1["A5"] = "۱. آمار تفکیکی وضعیت داوطلبان در مراحل ارزیابی (فایل مرجع اکسل)"
ws1["A5"].font = FONT_BOLD

stats_headers = ["مرحله ارزیابی", "کل داوطلبان", "مجاز / قبول", "غیرمجاز / مردود", "غایب (عدم حضور)", "درصد غیبت", "ملاحظات"]
for col_idx, h in enumerate(stats_headers, start=1):
    cell = ws1.cell(row=6, column=col_idx, value=h)
    cell.fill = NAVY_HEADER
    cell.font = FONT_HEADER
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER

stats_rows = [
    ["ثبت‌نام اولیه", 3377, 1428, 1949, 0, "۰.۰٪", "غربالگری مدارک و شرایط احراز سن/مدرک"],
    ["آزمون کتبی", 1157, 436, 479, 242, "۲۰.۹٪", "۲۴۲ داوطلب در جلسه آزمون کتبی حاضر نشدند"],
    ["آزمون مهارتی", 524, 218, 163, 143, "۲۷.۳٪", "۱۴۳ داوطلب در کانون آزمون عملی غیبت کردند"],
    ["مصاحبه تخصصی", 402, 144, 239, 19, "۴.۷٪", "۱۹ داوطلب در پنل مصاحبه حاضر نشدند"],
    ["کانون ارزیابی شایستگی", 155, 55, 100, 0, "۰.۰٪", "تمام ۱۵۵ نفر حاضر بودند (غایب نداشت)"],
    ["قبولی نهایی استخدامی", 49, 49, 0, 0, "۰.۰٪", "پذیرفته‌شدگان نهایی جهت طب صنعتی و بدو استخدام"],
    ["مجموع غیبت‌های ثبت شده", 404, "-", "-", 404, "-", "۴۰۴ مورد غیبت در کل فرآیند شناسایی گردید"]
]

for row_idx, r_data in enumerate(stats_rows, start=7):
    for col_idx, val in enumerate(r_data, start=1):
        cell = ws1.cell(row=row_idx, column=col_idx, value=val)
        cell.font = FONT_BOLD if row_idx == 13 else FONT_NORMAL
        cell.border = THIN_BORDER
        cell.alignment = ALIGN_CENTER if col_idx != 7 else ALIGN_RIGHT
        if row_idx == 13:
            cell.fill = AMBER_WARNING

# Table 2: Audit matching with ATS Database
ws1["A15"] = "۲. تطبیق وضعیت ۴۰۴ غایب شناسایی‌شده با پایگاه داده فعلی سامانه ATS"
ws1["A15"].font = FONT_BOLD

audit_headers = ["شرح انطباق با سامانه", "تعداد رکورد", "وضعیت فعلی در دیتابیس", "وضعیت پیشنهادی اصلاحی", "شیت مرجع جهت کنترل", "اقدام لازم توسط کارشناسان"]
for col_idx, h in enumerate(audit_headers, start=1):
    cell = ws1.cell(row=16, column=col_idx, value=h)
    cell.fill = STEEL_BLUE_SUB
    cell.font = FONT_HEADER
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER

audit_data = [
    ["غایبین دارای رکورد مردود در سامانه", 363, "مردود (FAILED)", "غایب (ABSENT)", "فهرست_کامل_۴۰۴_غایب", "تغییر وضعیت از FAILED به ABSENT در سامانه با نمره ۰.۰۰"],
    ["مغایرت بحرانی: غایب در اکسل، قبول در سامانه!", 2, "قبول (COMPLETED)", "نیاز به بررسی انسانی", "مغایرت‌ها_نیاز_به_تصمیم_انسان", "بررسی حامد ماجانی و محمدحسن اسماعیلی (کد شغل ۴۱۳۳)"],
    ["غایبین با مشاغل تعریف‌نشده در سامانه", 39, "فاقد رکورد مرحله", "ایجاد شغل و پرونده", "مغایرت‌ها_نیاز_به_تصمیم_انسان", "تعریف مشاغل ۴۱۲۰ (ایمنی) و ۴۱۲۹ (طرح و توسعه) و سپس ثبت غیبت"],
    ["جمع کل موارد غیبت", 404, "-", "۴۰۴ مورد غایب", "-", "کنترل ۱۰۰٪ توسط کارشناسان تأمین سرمایه انسانی"]
]

for row_idx, r_data in enumerate(audit_data, start=17):
    for col_idx, val in enumerate(r_data, start=1):
        cell = ws1.cell(row=row_idx, column=col_idx, value=val)
        cell.font = FONT_BOLD if row_idx == 20 else FONT_NORMAL
        cell.border = THIN_BORDER
        cell.alignment = ALIGN_CENTER if col_idx in [2, 3, 4] else ALIGN_RIGHT
        if row_idx == 18:
            cell.fill = RED_FAIL
        elif row_idx == 19:
            cell.fill = AMBER_WARNING
        elif row_idx == 20:
            cell.fill = ZEBRA_FILL

# Table 3: Guidelines for HR Team
ws1["A22"] = "۳. راهنمای گام‌به‌گام همکاران محترم جهت بررسی، کنترل و ورود اطلاعات به سامانه"
ws1["A22"].font = FONT_BOLD

steps = [
    ("گام ۱: بررسی مغایرت‌های بحرانی", "ابتدا به شیت «مغایرت‌ها_نیاز_به_تصمیم_انسان» مراجعه کرده و دو مورد تناقض ردیف‌های ۱ و ۲ (آقایان ماجانی و اسماعیلی) را بررسی فرمایید تا مشخص شود نمره موجود در سامانه صحیح است یا عدم حضور در اکسل."),
    ("گام ۲: ایجاد دو ردیف شغلی ۴۱۲۰ و ۴۱۲۹", "۳۹ داوطلب مربوط به مشاغل بازرس ایمنی (۴۱۲۰) و کارشناس کنترل پروژه (۴۱۲۹) هستند. این دو شغل را در سامانه ایجاد و متقاضیان را الصاق فرمایید."),
    ("گام ۳: کنترل فهرست ۴۰۴ غایب", "در شیت «فهرست_کامل_۴۰۴_غایب»، وضعیت تک‌تک غایبین با کد ملی و شناسه رکورد مشخص است. پس از تأیید انسانی، می‌توانید آن را به کار بگیرید."),
    ("گام ۴: ورود سریع از طریق شیت‌های آماده اکسل", "سه شیت «شیت_ورود_اکسل_کتبی»، «شیت_ورود_اکسل_مهارتی» و «شیت_ورود_اکسل_مصاحبه» دقیقاً مطابق فرمت سیستم آماده شده‌اند و می‌توان مستقیماً در بخش «ورود نمرات از اکسل» بارگذاری نمود."),
    ("گام ۵: حفظ سلامت و تفکیک داده‌ها", "با ثبت وضعیت «غایب»، سامانه به طور خودکار نمره را ۰.۰۰ قرار داده، مراحل بعدی را مسدود می‌کند و در داشبورد تحلیلی، نرخ غیبت تفکیک می‌گردد.")
]

for idx, (title, desc) in enumerate(steps, start=23):
    ws1.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=2)
    c_title = ws1.cell(row=idx, column=1, value=title)
    c_title.font = FONT_BOLD
    c_title.fill = ZEBRA_FILL
    c_title.border = THIN_BORDER
    c_title.alignment = ALIGN_RIGHT

    ws1.merge_cells(start_row=idx, start_column=3, end_row=idx, end_column=7)
    c_desc = ws1.cell(row=idx, column=3, value=desc)
    c_desc.font = FONT_NORMAL
    c_desc.border = THIN_BORDER
    c_desc.alignment = ALIGN_RIGHT

for col in range(1, 8):
    ws1.column_dimensions[get_column_letter(col)].width = 22
ws1.column_dimensions['A'].width = 25
ws1.column_dimensions['G'].width = 32

# -------------------------------------------------------------
# SHEET 2: مغایرت‌ها_نیاز_به_تصمیم_انسان (41 Items)
# -------------------------------------------------------------
ws2 = wb_out.create_sheet(title="مغایرت‌ها_نیاز_به_تصمیم_انسان")
format_sheet_view(ws2)

ws2.merge_cells("A1:M1")
ws2["A1"] = "فهرست موارد نیازمند کنترل و تصمیم‌گیری انسانی توسط کارشناسان جذب فولاد مبارکه (۴۱ مورد)"
ws2["A1"].font = FONT_TITLE
ws2["A1"].alignment = ALIGN_CENTER

headers_ws2 = [
    "ردیف", "کد ملی", "نام و نام خانوادگی", "کد شغل", "عنوان شغل", "واحد سازمانی",
    "مرحله ارزیابی", "وضعیت در اکسل", "وضعیت در سامانه", "نمره سامانه",
    "شرح مغایرت / مسئله", "پیشنهاد کارشناسی سامانه", "تصمیم و دستور کارشناس جذب (ثبت دستی)"
]

for col_idx, h in enumerate(headers_ws2, start=1):
    cell = ws2.cell(row=3, column=col_idx, value=h)
    cell.fill = NAVY_HEADER
    cell.font = FONT_HEADER
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER

ws2.freeze_panes = "A4"

for row_idx, item in enumerate(critical_discrepancies, start=4):
    rec = item['rec']
    cand = rec['candidate']
    cand_name = f"{cand.first_name} {cand.last_name}" if cand else "نامشخص در سامانه"
    j_title = rec['job'].title if rec['job'] else rec['job_info'].get('title', 'تعریف‌نشده')
    j_unit = rec['job'].unit if rec['job'] else rec['job_info'].get('unit', '-')
    ss = rec['stage_state']
    db_status = ss.get_status_display() if ss else "فاقد رکورد مرحله"
    db_score = ss.score if ss else "-"

    is_critical_pass = ss and ss.status == 'COMPLETED'

    row_vals = [
        row_idx - 3,
        rec['national_id'],
        cand_name,
        rec['exam_code'],
        j_title,
        j_unit,
        type_to_stage_name.get(rec['stage_type'], rec['stage_type']),
        "غایب (عدم حضور)",
        db_status,
        db_score,
        item['type'],
        item['action'],
        ""  # For human decision write-in
    ]

    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=val)
        cell.font = FONT_NORMAL
        cell.border = THIN_BORDER
        if col_idx in [1, 2, 4, 8, 9, 10]:
            cell.alignment = ALIGN_CENTER
        else:
            cell.alignment = ALIGN_RIGHT

        if is_critical_pass:
            cell.fill = RED_FAIL
        else:
            cell.fill = AMBER_WARNING if (row_idx % 2 == 0) else ZEBRA_FILL

# Auto width ws2
for col in range(1, 14):
    ws2.column_dimensions[get_column_letter(col)].width = 18
ws2.column_dimensions['A'].width = 8
ws2.column_dimensions['B'].width = 15
ws2.column_dimensions['C'].width = 22
ws2.column_dimensions['E'].width = 25
ws2.column_dimensions['K'].width = 35
ws2.column_dimensions['L'].width = 35
ws2.column_dimensions['M'].width = 30

# -------------------------------------------------------------
# SHEET 3: فهرست_کامل_۴۰۴_غایب
# -------------------------------------------------------------
ws3 = wb_out.create_sheet(title="فهرست_کامل_۴۰۴_غایب")
format_sheet_view(ws3)

ws3.merge_cells("A1:K1")
ws3["A1"] = "فهرست جامع ۴۰۴ داوطلب غایب در مراحل مختلف ارزیابی (جهت کنترل و اعمال در سامانه)"
ws3["A1"].font = FONT_TITLE
ws3["A1"].alignment = ALIGN_CENTER

headers_ws3 = [
    "ردیف", "شناسه وضعیت سامانه (State ID)", "کد ملی", "نام و نام خانوادگی", "کد شغل", "عنوان شغل",
    "مرحله ارزیابی", "وضعیت در فایل اکسل", "وضعیت فعلی سامانه", "وضعیت نهایی اصلاحی", "تأییدیه کارشناس جذب"
]

for col_idx, h in enumerate(headers_ws3, start=1):
    cell = ws3.cell(row=3, column=col_idx, value=h)
    cell.fill = NAVY_HEADER
    cell.font = FONT_HEADER
    cell.alignment = ALIGN_CENTER
    cell.border = THIN_BORDER

ws3.freeze_panes = "A4"

for row_idx, rec in enumerate(all_absentees, start=4):
    cand = rec['candidate']
    cand_name = f"{cand.first_name} {cand.last_name}" if cand else "نامشخص"
    j_title = rec['job'].title if rec['job'] else rec['job_info'].get('title', 'تعریف‌نشده')
    ss = rec['stage_state']
    state_id = ss.id if ss else "فاقد شناسه"
    db_status = ss.get_status_display() if ss else "فاقد رکورد"

    row_vals = [
        row_idx - 3,
        state_id,
        rec['national_id'],
        cand_name,
        rec['exam_code'],
        j_title,
        type_to_stage_name.get(rec['stage_type'], rec['stage_type']),
        "غایب",
        db_status,
        "غایب (ABSENT)",
        "تأیید شد [   ]"
    ]

    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws3.cell(row=row_idx, column=col_idx, value=val)
        cell.font = FONT_NORMAL
        cell.border = THIN_BORDER
        if col_idx in [1, 2, 3, 5, 7, 8, 9, 10, 11]:
            cell.alignment = ALIGN_CENTER
        else:
            cell.alignment = ALIGN_RIGHT

        if ss and ss.status == 'COMPLETED':
            cell.fill = RED_FAIL
        elif not ss:
            cell.fill = AMBER_WARNING
        else:
            cell.fill = GRAY_ABSENT if (row_idx % 2 == 0) else PatternFill(fill_type=None)

for col in range(1, 12):
    ws3.column_dimensions[get_column_letter(col)].width = 18
ws3.column_dimensions['A'].width = 8
ws3.column_dimensions['B'].width = 22
ws3.column_dimensions['C'].width = 16
ws3.column_dimensions['D'].width = 22
ws3.column_dimensions['F'].width = 28
ws3.column_dimensions['K'].width = 20

# -------------------------------------------------------------
# SHEETS 4, 5, 6: آماده ورود اکسل برای آزمون کتبی، مهارتی و مصاحبه
# Standard Import Format: [شناسه وضعیت, کد ملی, نام متقاضی, فرصت شغلی, نمره نهایی مرحله, وضعیت ارزیابی, توضیحات ارزیاب]
# -------------------------------------------------------------
stage_export_configs = [
    ("شیت_ورود_اکسل_کتبی", "کتبی", "آزمون کتبی"),
    ("شیت_ورود_اکسل_مهارتی", "مهارتی", "آزمون مهارتی"),
    ("شیت_ورود_اکسل_مصاحبه", "مصاحبه", "مصاحبه تخصصی")
]

for sheet_title, src_sheet, st_name in stage_export_configs:
    ws_imp = wb_out.create_sheet(title=sheet_title)
    format_sheet_view(ws_imp)

    headers_imp = ["شناسه وضعیت", "کد ملی", "نام متقاضی", "فرصت شغلی", "نمره نهایی مرحله", "وضعیت ارزیابی", "توضیحات ارزیاب"]
    for col_idx, h in enumerate(headers_imp, start=1):
        cell = ws_imp.cell(row=1, column=col_idx, value=h)
        cell.fill = NAVY_HEADER
        cell.font = FONT_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    ws_imp.freeze_panes = "A2"

    stage_records = excel_stage_data.get(src_sheet, [])
    for row_idx, rec in enumerate(stage_records, start=2):
        cand = rec['candidate']
        cand_name = f"{cand.first_name} {cand.last_name}" if cand else ""
        j_title = rec['job'].title if rec['job'] else rec['job_info'].get('title', '')
        ss = rec['stage_state']
        state_id = ss.id if ss else ""

        is_absent = 'غایب' in rec['result_raw']
        if is_absent:
            eval_status = "غایب"
            score_val = 0.0
            note_val = "غایب در این مرحله (ثبت تاریخی)"
        elif any(kw in rec['result_raw'] for kw in ['مجاز', 'قبول']):
            eval_status = "قبول"
            score_val = rec['score']
            note_val = "قبول مرحله"
        else:
            eval_status = "مردود"
            score_val = rec['score']
            note_val = "مردود علمی/مهارتی مرحله"

        row_vals = [
            state_id,
            rec['national_id'],
            cand_name,
            f"{rec['exam_code']} - {j_title}",
            score_val,
            eval_status,
            note_val
        ]

        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_imp.cell(row=row_idx, column=col_idx, value=val)
            cell.font = FONT_NORMAL
            cell.border = THIN_BORDER
            if col_idx in [1, 2, 5, 6]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_RIGHT

            if eval_status == "غایب":
                cell.fill = GRAY_ABSENT
            elif eval_status == "قبول":
                cell.fill = GREEN_PASS if (row_idx % 2 == 0) else PatternFill(fill_type=None)
            elif eval_status == "مردود":
                cell.fill = RED_FAIL if (row_idx % 2 == 0) else PatternFill(fill_type=None)

    for col in range(1, 8):
        ws_imp.column_dimensions[get_column_letter(col)].width = 20
    ws_imp.column_dimensions['A'].width = 16
    ws_imp.column_dimensions['B'].width = 16
    ws_imp.column_dimensions['C'].width = 22
    ws_imp.column_dimensions['D'].width = 30
    ws_imp.column_dimensions['G'].width = 30

# Remove default empty sheet
if default_sheet in wb_out.worksheets:
    wb_out.remove(default_sheet)

print("Saving workbook to workspace and downloads...")
wb_out.save(OUTPUT_EXCEL_WORKSPACE)
wb_out.save(OUTPUT_EXCEL_DOWNLOADS)

print("SUCCESS: Master audit workbook generated at:")
print(f"  - {OUTPUT_EXCEL_WORKSPACE}")
print(f"  - {OUTPUT_EXCEL_DOWNLOADS}")
